import os
from pathlib import Path
import time
from docx import Document
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel

load_dotenv(dotenv_path="../.env", override=True)
my_api_key=os.getenv("GROQ_API_KEY")


if not my_api_key:
    raise ValueError("API Key nahi h bro")

client=Groq(api_key=my_api_key)
model = os.getenv("MODEL")


job_description="""
BE/BTech/MCA or equivalent technical education.

Strong programming fundamentals and problem-solving ability.

Working knowledge of at least one language such as Python, JavaScript/TypeScript, Java or similar.

Basic understanding of APIs, databases, Git and Linux.

Curiosity about cloud, infrastructure, automation, Kubernetes and AI.

Ability to learn unfamiliar technologies independently and ask good technical questions.

Evidence that you build things: academic projects, GitHub work, internships, hackathons or personal projects.

What Matters More Than a Long Skills List

Fundamentals + curiosity + learning speed + ownership + ability to turn an idea into working software.
"""

class JobDescription(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience:float|None
    education_requirements:list[str]
    responsibilities:list[str]
    
jobD_schema = JobDescription.model_json_schema()

system_prompt = f""""
    You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{jobD_schema}

IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.

"""

user_prompt = f""""
Analyze the following job description
{job_description}
"""

message_system={
    "role":"system",
    "content":system_prompt
}

message_user={
    "role":"user",
    "content":user_prompt
}

response_format = {
    "type":"json_object"
}

messages=[message_system,message_user]

response=client.chat.completions.create(
    model=model,
    messages=messages,
    response_format=response_format
)
print("###################################")
answer=response.choices[0].message.content
# print(answer)


raw_json = answer

import json
job_data = json.loads(raw_json)

job = JobDescription(**job_data)

# print(job)


# parse resume

class MatchResult(BaseModel):
    score:float
    details:dict

class Experience(BaseModel):
    company:str|None = None
    role:str|None = None
    duration:str|None = None
    description:str|None = None
    skills_used:list[str]=[]

class Resume(BaseModel):
    name:str|None = None
    email:str|None = None
    phone:str|None = None
    
    total_experience_years : float|None = None
    
    skills:list[str]=[]
    experience : list[Experience]=[]
    education :list[str]=[]
    projects: list[str]=[]
    certifications : list[str]=[]
    
resume_schema = Resume.model_json_schema()


def parse_resume(resume_text):
    system_prompt = f"""
        You are an elite HR Technical Assistant and Information Extraction Engine.

        Your job is to parse candidate resumes and extract accurate, structured data.

        Parse the input resume text and extract candidate details based on semantic sections (e.g., "Experience", "Work History", "Internships", "Projects", "Skills").

        OUTPUT FORMAT:
        - Return ONLY a valid JSON object matching the exact schema provided.
        - Do NOT wrap the output in markdown code blocks (e.g., no ```json ... ```).
        - Match the structure and keys of this schema:

        {resume_schema}

        EXTRACTION RULES:
        1. **Fidelity:** Extract details verbatim or accurately summarized from the resume. Do NOT invent, assume, or fabricate any missing information.
        2. **Flexible Headings:** Identify relevant experience, projects, or education regardless of non-standard headings (e.g., "Employment History" vs "Work Experience").
        3. **Missing Data:** 
        - Return `null` for missing scalar fields (e.g., phone number, start dates).
        - Return an empty list `[]` for missing list fields.
        4. **Schema Compliance:** Output data values only. Do NOT include schema metadata keys like "properties", "type", or "description".
    """

    user_prompt = f""" 
    Parse the following resume:
    {resume_text}
    """
    message_system={
        "role":"system",
        "content":system_prompt
    }
    
    message_user={
        "role":"user",
        "content":user_prompt
    }
    
    messages = [message_system,message_user]
    response_format={
        "type":"json_object"
    }
    
    response = client.chat.completions.create(
        model = model,
        messages=messages,
        response_format=response_format
    )
    
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
        You are an expert HR Technical Recruiter and Candidate Evaluator.

        Your task is to conduct a detailed, objective comparison between a candidate's resume and a job description.

        JOB DESCRIPTION:
        {job.model_dump_json(indent=2)}

        CANDIDATE RESUME:
        {resume.model_dump_json(indent=2)}

        OUTPUT REQUIREMENTS:
        - Return ONLY a valid JSON object matching the schema below.
        - Do NOT wrap the JSON in markdown code blocks (e.g., no ```json ... ```).
        - Output must conform exactly to this schema:

        {match_schema}

        EVALUATION CRITERIA:
        1. **Candidate Name:** Extract the primary name from the candidate's resume.
        2. **Matching Skills:** List key skills present in both the resume and job description (accounting for synonyms or equivalent tech stacks).
        3. **Missing Important Skills:** List critical required skills specified in the job description that are missing or weakly demonstrated in the resume.
        4. **Experience Requirement:** Evaluate total relevant experience against the minimum required years and return a boolean/status indicating whether the requirement is met.
        5. **Match Percentage:** Calculate an overall score from 0 to 100 based on skill overlap, experience depth, and core qualifications.
        6. **Final Verdict:** Provide a 1 to 2 sentence executive summary explaining the rating and immediate next step (e.g., "Strong fit for interview", "Lacks core experience").
    """

    message = {
        "role":"user",
        "content":prompt
    }
    messages = [message]
    response_format={
        "type":"json_object"
    }
    response = client.chat.completions.create(
        model=model,
        messages = messages,
        response_format=response_format
    )
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

# resume parse
from pypdf import PdfReader
from docx import Document

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text=""
    for page in reader.pages:
        page_text = page.extract_text()
        if(page_text):
            text+=page_text+"\n"
    return text

def read_docx(file_path):
    doc = Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text+=paragraph.text + "\n"
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text+=cell.text+"\n"
    
    return text

def read_resume(file_path):
    if file_path.suffix.lower()=='.pdf':
        return read_pdf(file_path)
    elif file_path.suffix.lower()==".docx":
        return read_docx(file_path)
    else:
        return None
    
# 
resume_folder = Path("resumes")
all_results=[]

for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue
    
    print("\n Processing: ",file_path.name)
    
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)
    
    time.sleep(5)
    
    result = final_score(job,parsed_resume)
    
    time.sleep(5)
    
    print("Score:",result.score)
    all_results.append({
            "name":parsed_resume.name,
            "score":result.score,
            "details":result.details,
        }  
    )

all_results.sort(
    key=lambda candidate:candidate["score"],
    reverse=True
)

top_2 = all_results[:2]

print("Top 2 candidates")

for candidate in top_2:
    
    print(
        candidate["name"],"-"
        ,candidate["score"],"%"
    )
    
    print(candidate["details"])
    
